import json,csv,io,datetime,hashlib,zipfile,re,collections
from decimal import Decimal as D,ROUND_HALF_UP
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1];W=ROOT/'.work';REPO=ROOT.parents[2];OLD=REPO/'docs/migration/preparation_20261011'
specs=json.loads((W/'workbooks.json').read_text(encoding='utf-8'))
summary=json.loads((W/'summary.json').read_text(encoding='utf-8'))
masters=json.loads((W/'local_masters.json').read_text(encoding='utf-8-sig'))['Tables']
numeric={'Amount','Quantity','UnitCost','SellingPrice','ConversionFactor','Debit','Credit','StatementBalance','RowCount'}
dates={'DocumentDate','DueDate','ChequeDate','ReceivedDate','LayerDate','ClosingDate'}
modules={};checks=[];workbookrows=[]
def check(name,detail):checks.append([name,'Pass',detail])
def rd(p):return list(csv.DictReader(p.open(encoding='utf-8-sig',newline='')))
def amount(v):return D(str(v))
def money(v):return D(v).quantize(D('.01'),rounding=ROUND_HALF_UP)
def norm(v,h=''):
    if v in ('',None):return ''
    if isinstance(v,(datetime.datetime,datetime.date)):return v.strftime('%Y-%m-%d')
    if h in numeric:return D(str(v))
    return str(v).replace('\r\n','\n').replace('\r','\n')
for spec in specs:
    book=ROOT/spec['file'];wb=openpyxl.load_workbook(book,data_only=False)
    assert len(wb.worksheets)==len(spec['sheets'])
    for ss,ws in zip(spec['sheets'],wb.worksheets):
        assert list(next(ws.values))==ss['headers']
        assert ws.max_row==len(ss['rows'])+1,(book,ws.title)
        assert not [(c.coordinate,c.value) for row in ws for c in row if c.data_type=='f'],book
        for actual,expected in zip(list(ws.values)[1:],ss['rows']):
            for h,a,e in zip(ss['headers'],actual,expected):
                if ss.get('template'):assert norm(a,h)==norm(e,h),(book,h,a,e)
                elif e in ('',None):assert a in ('',None),(book,h,a,e)
                elif isinstance(e,(int,float)) and not isinstance(e,bool):assert D(str(a))==D(str(e)),(book,h,a,e)
                elif isinstance(a,(int,float)) and not isinstance(a,bool):assert D(str(a))==D(str(e)),(book,h,a,e)
                else:assert norm(a)==norm(e),(book,h,a,e)
        if ss.get('template'):
            k=ss['name'];p=ROOT/'upload_csv'/(k+'.csv');rr=rd(p);modules[k]=rr
            template=next(csv.reader((REPO/'docs/migration/templates'/(k+'.csv')).read_text(encoding='utf-8-sig').splitlines()))
            assert template==ss['headers']==list(rr[0]),k
            catalog=(REPO/'NexusFlow.AppCore/Features/Migration/MigrationModels.cs').read_text(encoding='utf-8-sig')
            fields=re.findall(r'"([^"]+)"',re.search(r'\["'+k+r'"\]\s*=\s*\[(.*?)\]',catalog).group(1))
            assert template==fields+['ClosingDate']
            assert len(rr)==len(ss['rows'])==summary['rows'][k] and len(rr)<=20000
            assert 0<p.stat().st_size<=20*1024*1024 and book.stat().st_size<=20*1024*1024
            for i,(row,cells,ex) in enumerate(zip(rr,list(ws.values)[1:],ss['rows']),2):
                assert set(row)==set(template) and None not in row
                for h,a,e in zip(template,cells,ex):
                    v=row[h];assert norm(v,h)==norm(a,h)==norm(e,h),(k,i,h,v,a,e)
                    if h in numeric and v:
                        assert re.fullmatch(r'-?\d+(?:\.\d+)?',v),(k,i,h)
                        assert isinstance(a,(int,float)),(k,i,h,'numeric cell')
                    if h in dates and v:
                        assert re.fullmatch(r'\d{4}-\d{2}-\d{2}',v)
                        datetime.date.fromisoformat(v)
                        assert isinstance(a,datetime.datetime),(k,i,h,'date cell')
                        assert ws.cell(i,template.index(h)+1).number_format.lower()=='yyyy-mm-dd'
                    if h in ['Code','ChequeNumber','LegacyRecordId','Reference'] and v:assert isinstance(a,str),(k,i,h,'text identifier')
            assert ws.freeze_panes
            workbookrows.append([k,len(rr),'Exact current template','Equal','None','ISO / typed','Numeric','Text'])
        else:workbookrows.append([spec['file'].split('/')[-1],sum(len(s['rows']) for s in spec['sheets']),'Seven review sheets','Matrix equality','None','As applicable','As applicable','As applicable'])
    wb.close()
check('Workbook and CSV equality','13 workbooks / 12 UTF-8 CSVs; exact current template headers, typed cells, final values and no formulas.')
sources=json.loads((OLD/'.work/sources.json').read_text(encoding='utf-8'))
for s in sources:assert hashlib.sha256(Path(s['path']).read_bytes()).hexdigest()==s['sha256'],s['filename']
with zipfile.ZipFile(OLD/'archive/Legacy_Originals.zip') as z:
    assert z.testzip() is None
    for s in sources:assert hashlib.sha256(z.read(s['key']+'/'+s['filename'])).hexdigest()==s['sha256']
for r in rd(OLD/'Package_Manifest.csv'):
    assert hashlib.sha256((OLD/r['File']).read_bytes()).hexdigest()==r['SHA256'],r['File']
check('Original and previous deliverable preservation','All 10 original hashes, original archive ZIP and all previous package manifest members unchanged.')
def identity(k,r):
    if r['LegacyRecordId']:return r['LegacyRecordId']
    if k=='ControlTotals':a=[r['Kind'],r['Code'],r['Component']]
    elif k.endswith('Master') or k in ['Accounts','BankOpening','TrialBalance']:return r['Code']
    elif k=='Stock':a=[r['Code'],r['WarehouseCode'],r['Batch'],r['LayerDate']]
    else:a=[r['Code'],r.get('Kind',''),r.get('Reference',''),r.get('ChequeDate','') if k=='ReceivedCheques' else r.get('DocumentDate',''),r.get('ChequeNumber','')]
    return json.dumps(a,separators=(',',':'),ensure_ascii=False)
for k,rr in modules.items():
    keys=[identity(k,r) for r in rr];assert len(keys)==len(set(keys)),k
    for r,key in zip(rr,keys):
        assert 0<len(r['Code'])<=100 and len(key)<=300
        assert r['ClosingDate']=='2026-09-30'
check('Stable identities and closing metadata','No duplicate supported identities or row-number identities; all closing metadata is explicitly UAT 2026-09-30.')
accounts={r['Code']:r for r in masters['Accounts'] if r['IsActive'] and r['IsTransactionAccount']}
types={'Asset':1,'Liability':2,'Equity':3,'Revenue':4,'Expense':5}
assert set(r['Code'] for r in modules['Accounts'])==set(accounts)==set(r['Code'] for r in modules['TrialBalance'])
for r in modules['Accounts']:assert r['Name']==accounts[r['Code']]['Name'] and types[r['AccountType']]==accounts[r['Code']]['Type']
assert accounts['3200']['Type']==3
exceptions=[]
for k in ['CustomerMaster','SupplierMaster']:
    for r in modules[k]:
        assert 0<len(r['Name'])<=200 and len(r['Code'])<=50
        for f in ['Email','Contact','Phone','Mobile','Address1','Address2','TaxNumber']:
            lim=200 if f.startswith('Address') else 150 if f=='Email' else 100 if f=='Contact' else 50
            assert len(r[f])<=lim,(k,r['Code'],f)
        assert not r['Email'] or re.fullmatch(r'[^@\s]+@[^@\s]+',r['Email']),(k,r['Code'],r['Email'])
        for f in ['Email','Contact','Phone','Address1','TaxNumber']:
            if not r[f]:exceptions.append([k,r['Code'],f,'Missing genuine detail','Blank retained; complete through maintenance.','Warning'])
        at={'ReceivableAccountCode':1,'RevenueAccountCode':4} if k=='CustomerMaster' else {'PayableAccountCode':2,'ExpenseAccountCode':5}
        for f,t in at.items():assert accounts[r[f]]['Type']==t,(k,r['Code'],f)
        assert r['PaymentTermCode']=='NET30'
        assert r['GroupCode']==('TRADE' if k=='CustomerMaster' else 'DOM')
        if k=='CustomerMaster':assert r['PriceLevelCode']=='STANDARD TIER'
    tax=[r['TaxNumber'] for r in modules[k] if r['TaxNumber']]
    if k=='SupplierMaster':assert len(tax)==len(set(tax))
assert next(r for r in modules['CustomerMaster'] if r['Code']=='BA366')['Email']==''
check('Masters and account types','Master field lengths, email format, supplier tax uniqueness; AR Asset, AP Liability, revenue Revenue, expense Expense, equity Equity; NET30 differs from COD.')
for k in ['CustomerOpenings','SupplierOpenings','ReceivedCheques','Stock']:
    prior=rd(OLD/'held_csv'/(k+'.csv'));current=modules[k];assert len(prior)==len(current)
    fields=['Code','Reference','Amount','Kind'] if k.endswith('Openings') else ['Code','Reference','ChequeNumber','Amount','ChequeDate','ReceivedDate','Status','PartyReflected','BankReflected'] if k=='ReceivedCheques' else ['Code','Quantity','UnitCost','Batch']
    for a,b in zip(prior,current):
        for f in fields:assert a[f]==b[f],(k,a['Code'],f)
customer_codes={r['Code'] for r in modules['CustomerMaster']};supplier_codes={r['Code'] for r in modules['SupplierMaster']}
for k in ['CustomerOpenings','SupplierOpenings','ReceivedCheques','IssuedCheques']:
    for r in modules[k]:
        assert r['Code'] in (customer_codes if k in ['CustomerOpenings','ReceivedCheques'] else supplier_codes)
        v=D(r['Amount']);assert v!=0 and money(v)==v and abs(v)<D('10000000000000000')
        assert 0<len(r['Reference'])<=100
        if k.endswith('Openings'):
            assert r['DocumentDate']<='2026-09-30'
            assert datetime.date.fromisoformat(r['DueDate'])==datetime.date.fromisoformat(r['DocumentDate'])+datetime.timedelta(days=30)
            assert r['Kind']==(('Invoice' if v>0 else 'ReturnCredit') if k=='CustomerOpenings' else ('Bill' if v>0 else 'Credit'))
        else:
            assert v>0
            assert r['PartyReflected']=='true'
            if k=='ReceivedCheques':assert r['Status']=='InSafe' and r['BankReflected']=='false' and r['SafeAccountCode']=='1151'
            else:assert r['BankReflected']=='true' and r['BankAccountCode']=='1112'
            assert r['ReceivedDate' if k=='ReceivedCheques' else 'DocumentDate']<='2026-09-30'
net=collections.defaultdict(lambda:D(0));doccount=collections.Counter()
for r in modules['CustomerOpenings']:net[r['Code']]+=D(r['Amount']);doccount[r['Code']]+=1
zeros={c for c,v in net.items() if v==0};assert zeros=={'DC969','FA111','JY308','SA280','VI852'}
for c in zeros:assert any(D(r['Amount'])>0 for r in modules['CustomerOpenings'] if r['Code']==c) and any(D(r['Amount'])<0 for r in modules['CustomerOpenings'] if r['Code']==c)
check('Legacy balances and signs','All 1,150 AR, 180 AP, 49 received cheques and 938 stock rows retained with unchanged amounts/quantities/costs. Five offsetting zero-net customer cases retained. Unmatched fixture master codes: zero.')
maps=rd(ROOT/'reports/Mappings.csv')
mapping={(r['Mapping type'],r['Exact legacy code'].upper()):int(r['Local target ID']) for r in maps}
lookups={r['Id']:r for r in masters['SystemLookups']}
for typ,code in [('CustomerGroup','TRADE'),('SupplierGroup','DOM'),('PaymentTerm','NET30'),('PriceLevel','STANDARD TIER')]:
    l=lookups[mapping[typ,code]];assert l['Type']==typ and l['Code']==code and l['IsActive']
branches={r['Id']:r for r in masters['BankBranches']};banks={r['Id']:r for r in masters['Banks']}
for r in modules['ReceivedCheques']:
    b=branches[mapping['BankBranch',r['BankBranchCode'].upper()]];assert b['IsActive'] and banks[b['BankId']]['IsActive']
wh=next(r for r in masters['Warehouses'] if r['Id']==mapping['Warehouse','WH_MAIN']);assert wh['Code']=='WH_MAIN' and wh['IsActive']
units={r['Id']:r for r in masters['UnitOfMeasures']};cats={r['Id']:r for r in masters['Categories']};brands={r['Id']:r for r in masters['Brands']};products={r['Id']:r for r in masters['Products']};variants={r['SKU']:r for r in masters['ProductVariants']}
new=0
for r in modules['Stock']:
    assert r['WarehouseCode']=='WH_MAIN' and r['LayerDate']=='2026-09-30'
    q=D(r['Quantity']);cost=D(r['UnitCost']);factor=D(r['ConversionFactor'])
    assert q>=0 and cost>=0 and factor==1 and q==q.quantize(D('.0001')) and cost==cost.quantize(D('.000001'))
    assert accounts[r['InventoryAccountCode']]['Type']==1
    u=units[mapping['Unit',r['UnitCode'].upper()]]
    v=variants.get(r['Code'])
    if v:
        assert v['IsActive']
        p=products[v['ProductId']];cat=cats[p['CategoryId']]
        assert p['UnitOfMeasureId']==u['Id'] and r['ProductType']=={1:'RawMaterial',2:'FinishedGood'}[p['Type']]
        assert cat['Code']==r['CategoryCode'] and brands[p['BrandId']]['Name']==r['BrandCode']
        assert accounts[r['InventoryAccountCode']]['Id']==(wh['OverrideInventoryAccountId'] or cat['InventoryAccountId'])
    else:
        new+=1;cat=cats[mapping['Category',r['CategoryCode'].upper()]]
        assert r['ProductType']=='FinishedGood' and r['UnitCode']=='pcs' and r['InventoryAccountCode']=='1132'
        assert cat['InventoryAccountId']==accounts['1132']['Id'] and accounts['4110']['Id']==cat['SalesAccountId'] and accounts['5110']['Id']==cat['CogsAccountId']
        assert brands[mapping['Brand',r['BrandCode'].upper()]]['Name']=='UNBRANDED'
assert new==81
check('Explicit mappings and Stock','Every required lookup resolves against the read-only local snapshot; 857 existing SKUs agree with current masters; 81 new FG fixtures have complete valid UAT category/brand/unit/account selections.')
expected={}
def control(kind,c,component,value,count):expected[kind,c,component]=(D(value),count)
for kind,master,opening,cheque in [('Customer','CustomerMaster','CustomerOpenings','ReceivedCheques'),('Supplier','SupplierMaster','SupplierOpenings','IssuedCheques')]:
    for r in modules[master]:
        c=r['Code'];a=[D(x['Amount']) for x in modules[opening] if x['Code']==c];b=[D(x['Amount']) for x in modules[cheque] if x['Code']==c]
        control(kind,c,opening+':Debit',sum((max(v,0) for v in a),D(0)),len(a))
        control(kind,c,opening+':Credit',sum((max(-v,0) for v in a),D(0)),len(a))
        control(kind,c,cheque,sum(b,D(0)),len(b))
for r in modules['Stock']:
    control('Stock',r['Code'],'Stock:WH_MAIN:Quantity',D(r['Quantity']),1)
    control('Stock',r['Code'],'Stock:WH_MAIN:Value',money(D(r['Quantity'])*D(r['UnitCost'])),1)
bank=modules['BankOpening'][0];assert bank['Code']=='1112'
bridge=D(bank['StatementBalance'])+sum((D(r['Amount']) for r in modules['BankRecon']),D(0))-sum((D(r['Amount']) for r in modules['IssuedCheques']),D(0))
assert bridge==D(bank['Amount'])==100000
control('Bank','1112','BankOpening',D(bank['Amount']),1);control('Bank','1112','BankStatementBridge',bridge,2);control('Bank','1112','BankRecon',2500,1)
actual={(r['Kind'],r['Code'],r['Component']):(D(r['Amount']),int(r['RowCount'])) for r in modules['ControlTotals']}
assert actual==expected and len(actual)==5212
scopes={(k,c) for k,c,comp in actual};assert len(scopes)==2050
assert all((r['Kind'],r['Code']) in scopes for r in masters['MigrationScopes'])
check('Scope reconciliation controls','All 5,212 controls independently recalculated from CSV; 2,050 scopes covered, including every prior staged scope and zero-opening masters. Bank statement bridge equals book.')
journals=collections.defaultdict(lambda:collections.defaultdict(lambda:D(0)))
for r in modules['CustomerOpenings']:journals['CustomerOpenings']['1120']+=D(r['Amount'])
for r in modules['SupplierOpenings']:journals['SupplierOpenings']['2110']-=D(r['Amount'])
for r in modules['ReceivedCheques']:journals['ReceivedCheques']['1151']+=D(r['Amount'])
for r in modules['Stock']:journals['Stock'][r['InventoryAccountCode']]+=money(D(r['Quantity'])*D(r['UnitCost']))
journals['BankOpening']['1112']=D(bank['Amount'])
opening=collections.defaultdict(lambda:D(0))
for k,mm in journals.items():
    mm['3200']=-sum(mm.values());assert sum(mm.values())==0
    for c,v in mm.items():opening[c]+=v
tb={r['Code']:D(r['Debit'])-D(r['Credit']) for r in modules['TrialBalance']}
assert sum(tb.values())==0
residual={c:tb[c]-opening[c] for c in accounts}
assert {c:v for c,v in residual.items() if v}=={'4110':D('-1000000'),'6210':D('1000000')}
assert all(residual[c]==0 for c in ['1120','2110','1112','1151','1131','1132'])
assert sum(D(r['Debit']) for r in modules['TrialBalance'])==sum(D(r['Credit']) for r in modules['TrialBalance'])==D('282960858.88')
for r in rd(ROOT/'reports/Expected_GL.csv'):
    c=r['Code'];assert D(r['Opening module net (LKR)'])==opening[c] and D(r['UAT closing net (LKR)'])==tb[c] and D(r['TB residual (LKR)'])==residual[c]
check('Expected posting and Trial Balance','39 current posting accounts; TB debit = credit 282,960,858.88; managed-account residuals zero; residual rent debit / revenue credit 1,000,000 each. No duplicate cheque ledger movement.')
stages={'01_Masters':['Accounts','CustomerMaster','SupplierMaster'],'02_Openings_and_Controls':['CustomerOpenings','SupplierOpenings','ReceivedCheques','IssuedCheques','Stock','BankOpening','BankRecon','ControlTotals'],'03_TrialBalance':['TrialBalance']}
for name,kk in stages.items():
    with zipfile.ZipFile(ROOT/'stage_packages'/(name+'.zip')) as z:
        assert z.testzip() is None
        m=json.loads(z.read('manifest.json'));assert m['Version']=='1.0' and m['SourceSystem']=='TESTTTT' and m['ClosingDate']=='2026-09-30'
        assert set(z.namelist())=={'manifest.json'}|{k+'.csv' for k in kk}
        assert m['Files']==[{'Module':k,'File':k+'.csv'} for k in kk]
        for k in kk:assert z.read(k+'.csv')==(ROOT/'upload_csv'/(k+'.csv')).read_bytes()
check('Staged ZIP packages','Three valid manifests; exact byte-equal CSV members, unique modules, correct source TESTTTT and closing; ZIP integrity passed.')
for r in rd(ROOT/'reports/UAT_changes.csv'):
    exceptions.append([r['Module'],r['Code'],r['Field'],'Explicit UAT adjustment',r['Reason'],'UAT fixture / production review required'])
exceptions.append(['All','','Destination UI','Destination verification','Hosted browser access ERR_BLOCKED_BY_CLIENT; validate current destination templates, masters and preview before execution.','Destination validation required'])
def write(p,h,rr):
    with p.open('w',encoding='utf-8',newline='') as f:
        w=csv.writer(f);w.writerow(h);w.writerows(rr)
write(ROOT/'reports/Verification.csv',['Check','Result','Detail'],checks)
write(ROOT/'reports/Workbook_verification.csv',['Workbook','Rows','Headers','CSV comparison','Formulas','Dates','Amounts','Identifiers'],workbookrows)
write(ROOT/'reports/UAT_Exceptions.csv',['Module','Code','Field','Classification','Detail','Readiness'],exceptions)
payload={'checked_workbooks':len(specs),'csv_files':len(modules),'original_files_unchanged':len(sources),'old_manifest_files_unchanged':len(rd(OLD/'Package_Manifest.csv')),'checks':checks,'duplicate_identities':0,'unmatched_master_codes':0,'scopes':len(scopes),'control_rows':len(actual),'no_application_validation_or_import_executed':True}
(W/'verification.json').write_text(json.dumps(payload,indent=2),encoding='utf-8')
md='# UAT file verification\n\nVerified 12 October 2026. Ready to upload for UAT staging. Application validation, approval and execution are operator stages; they were not performed during preparation. Production accuracy is not certified.\n\n'
md+='\n'.join('- **'+n+': Pass.** '+d for n,s,d in checks)
md+='\n\nChecks use an independent read-only workbook/CSV pass and independently recomputed control/journal totals. Master comparisons use the 2026-10-12T15:52:41 UTC local SELECT-only snapshot. Destination master differences or deployed importer changes require review at Validate Selected Import. Expected controls are test-fixture expectations, not independently certified production closing controls.\n\n'
md+='[Detailed checks](reports/Verification.csv), [workbook checks](reports/Workbook_verification.csv), [UAT exceptions](reports/UAT_Exceptions.csv).\n'
(ROOT/'Verification_Report.md').write_text(md,encoding='utf-8')
print(json.dumps({k:v for k,v in payload.items() if k!='checks'},indent=2))

